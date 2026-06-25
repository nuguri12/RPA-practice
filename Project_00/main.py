#!/usr/bin/env python3
import argparse
import os
import re
import smtplib
import ssl
import sys
from email.message import EmailMessage
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl이 설치되지 않았습니다. 'pip install -r requirements.txt' 를 먼저 실행하세요.", file=sys.stderr)
    sys.exit(1)


def normalize_header(value):
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9가-힣]+", "", str(value).strip().lower())


def load_employee_email_map(db_path):
    wb = load_workbook(db_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    if not rows:
        wb.close()
        raise ValueError("엑셀 파일에 데이터가 없습니다.")

    headers = [normalize_header(cell) for cell in rows[0]]

    id_col = None
    email_col = None

    for idx, header in enumerate(headers):
        if header in {"employeeid", "employee_id", "사번", "empid", "id", "employeeidnumber"}:
            id_col = idx
        if header in {"email", "e-mail", "emailaddress", "메일", "mail", "emailaddr"}:
            email_col = idx

    # 헤더를 못 찾으면 첫째/둘째 열로 가정
    if id_col is None or email_col is None:
        if len(headers) >= 2:
            id_col, email_col = 0, 1
        else:
            wb.close()
            raise ValueError("엑셀 DB에서 사번/이메일 열을 찾지 못했습니다.")

    mapping = {}
    for row in rows[1:]:
        if not row:
            continue
        emp_id = str(row[id_col] or "").strip()
        email = str(row[email_col] or "").strip()
        if emp_id and email:
            mapping[emp_id] = email

    wb.close()
    return mapping


def extract_employee_id(file_name):
    stem = Path(file_name).stem
    m = re.match(r"^([0-9A-Za-z]+)", stem)
    return m.group(1) if m else None


def discover_files(folder):
    files = []
    for root, _, filenames in os.walk(folder):
        for name in sorted(filenames):
            path = Path(root) / name
            if path.is_file():
                files.append(path)
    return files


def send_mail(host, port, user, password, sender_email, recipient, subject, body, attachment_path, dry_run):
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = sender_email
    msg["To"] = recipient
    msg.set_content(body)

    with attachment_path.open("rb") as f:
        msg.add_attachment(
            f.read(),
            maintype="application",
            subtype="octet-stream",
            filename=attachment_path.name,
        )

    if dry_run:
        print(f"[DRY RUN] {attachment_path.name} -> {recipient}")
        return

    context = ssl.create_default_context()
    with smtplib.SMTP(host, port, timeout=30) as server:
        server.starttls(context=context)
        server.login(user, password)
        server.send_message(msg)

    print(f"발송 완료: {attachment_path.name} -> {recipient}")


def main():
    parser = argparse.ArgumentParser(description="사번 기반 급여명세서 자동 메일 발송")
    parser.add_argument("--folder", required=True, help="급여명세서 파일이 있는 폴더")
    parser.add_argument("--db", required=True, help="사번/이메일이 있는 엑셀 파일")
    parser.add_argument("--smtp-host", required=True, help="SMTP 서버 주소")
    parser.add_argument("--smtp-port", type=int, default=587, help="SMTP 포트")
    parser.add_argument("--smtp-user", required=True, help="SMTP 계정")
    parser.add_argument("--smtp-password", required=True, help="SMTP 비밀번호")
    parser.add_argument("--sender-email", required=True, help="보내는 사람 메일")
    parser.add_argument("--subject", default="급여명세서", help="메일 제목")
    parser.add_argument("--message", default="급여명세서를 첨부하였습니다.", help="메일 본문")
    parser.add_argument("--dry-run", action="store_true", help="실제 발송 없이 실행만 확인")
    args = parser.parse_args()

    folder = Path(args.folder).expanduser()
    db_path = Path(args.db).expanduser()

    if not folder.exists() or not folder.is_dir():
        raise SystemExit(f"폴더를 찾을 수 없습니다: {folder}")

    if not db_path.exists():
        raise SystemExit(f"엑셀 DB를 찾을 수 없습니다: {db_path}")

    employee_map = load_employee_email_map(str(db_path))

    for path in discover_files(folder):
        emp_id = extract_employee_id(path.name)
        if not emp_id:
            continue

        recipient = employee_map.get(emp_id)
        if not recipient:
            print(f"미매칭: {path.name} (사번={emp_id})")
            continue

        send_mail(
            host=args.smtp_host,
            port=args.smtp_port,
            user=args.smtp_user,
            password=args.smtp_password,
            sender_email=args.sender_email,
            recipient=recipient,
            subject=args.subject,
            body=args.message,
            attachment_path=path,
            dry_run=args.dry_run,
        )


if __name__ == "__main__":
    main()